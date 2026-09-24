"""Построение журнала посещаемости за месяц и экспорт в CSV / Excel."""
import calendar
import csv
import datetime
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from django.utils.formats import date_format
from django.utils.translation import gettext as _

from core.models import Attendance, attendance_percent


def parse_month(value):
    """'2026-09' → date(2026, 9, 1); при ошибке — текущий месяц."""
    try:
        year, month = (int(p) for p in (value or "").split("-")[:2])
        return datetime.date(year, month, 1)
    except (ValueError, TypeError):
        today = timezone.localdate()
        return today.replace(day=1)


def shift_month(first_day, delta):
    month = first_day.month - 1 + delta
    return datetime.date(first_day.year + month // 12, month % 12 + 1, 1)


def build_month_journal(group, month_start):
    last_day = month_start.replace(day=calendar.monthrange(month_start.year, month_start.month)[1])
    records = list(
        Attendance.objects.filter(group=group, date__range=(month_start, last_day)).select_related("student")
    )
    dates = sorted({r.date for r in records} | set(group.lesson_dates_in_range(month_start, last_day)))

    students = list(group.students.order_by("last_name", "first_name"))
    known = {s.pk for s in students}
    # Студенты, которые ушли из группы, но имеют отметки в этом месяце, тоже остаются в журнале.
    for r in records:
        if r.student_id not in known:
            students.append(r.student)
            known.add(r.student_id)

    by_key = {(r.student_id, r.date): r for r in records}
    rows = []
    for student in students:
        cells = [by_key.get((student.pk, d)) for d in dates]
        marked = [c for c in cells if c]
        grades = [c.grade for c in marked if c.grade]
        rows.append(
            {
                "student": student,
                "cells": cells,
                "slots": list(zip(dates, cells)),
                "percent": attendance_percent(marked),
                "avg_grade": round(sum(grades) / len(grades), 1) if grades else None,
                "counts": {s: sum(1 for c in marked if c.status == s) for s in Attendance.Status.values},
            }
        )

    per_date = []
    for i, d in enumerate(dates):
        marked = [row["cells"][i] for row in rows if row["cells"][i]]
        per_date.append({"date": d, "percent": attendance_percent(marked), "marked": bool(marked)})

    return {
        "group": group,
        "month": month_start,
        "last_day": last_day,
        "prev_month": shift_month(month_start, -1),
        "next_month": shift_month(month_start, 1),
        "is_current_month": month_start == timezone.localdate().replace(day=1),
        "dates": dates,
        "per_date": per_date,
        "rows": rows,
        "percent": attendance_percent(records),
        "total_marks": len(records),
    }


def _cell_text(record):
    if not record:
        return ""
    text = record.get_status_display()
    if record.grade:
        text += f" ({record.grade})"
    return text


def _export_rows(journal):
    header = [_("Student")] + [d.strftime("%d.%m") for d in journal["dates"]] + [_("Attendance, %"), _("Average grade")]
    rows = []
    for row in journal["rows"]:
        rows.append(
            [row["student"].display_name]
            + [_cell_text(c) for c in row["cells"]]
            + ["" if row["percent"] is None else row["percent"], row["avg_grade"] or ""]
        )
    footer = [_("Group")] + ["" if p["percent"] is None else f'{p["percent"]}%' for p in journal["per_date"]]
    footer += ["" if journal["percent"] is None else journal["percent"], ""]
    return header, rows, footer


def _filename(journal, ext):
    return f'journal_{journal["group"].pk}_{journal["month"]:%Y-%m}.{ext}'


def export_csv(journal):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(journal, "csv")}"'
    response.write("﻿")  # BOM — чтобы Excel корректно открыл кириллицу
    writer = csv.writer(response, delimiter=";")
    header, rows, footer = _export_rows(journal)
    writer.writerow([f'{journal["group"].name} — {date_format(journal["month"], "F Y")}'])
    writer.writerow(header)
    writer.writerows(rows)
    writer.writerow(footer)
    return response


STATUS_FILL = {
    Attendance.Status.PRESENT: "DCFCE7",
    Attendance.Status.LATE: "FEF3C7",
    Attendance.Status.ABSENT: "FEE2E2",
    Attendance.Status.EXCUSED: "DBEAFE",
}


def export_xlsx(journal):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f'{journal["month"]:%Y-%m}'
    header, rows, footer = _export_rows(journal)

    ws.append([f'{journal["group"].name} — {date_format(journal["month"], "F Y")}'])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(header)
    thin = Side(style="thin", color="D4D4D8")
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A0A0A")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for data_row, row in zip(rows, journal["rows"]):
        ws.append(data_row)
        excel_row = ws.max_row
        for i, record in enumerate(row["cells"], start=2):
            cell = ws.cell(row=excel_row, column=i)
            cell.alignment = Alignment(horizontal="center")
            if record:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL[record.status])
    ws.append(footer)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.column_dimensions["A"].width = 28
    for col in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B3"

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(journal, "xlsx")}"'
    return response
